from __future__ import annotations

import re
import os
from pathlib import Path

import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    import matplotlib.pyplot as plt
except ModuleNotFoundError:  # Bundled document runtime omits plotting extras.
    plt = None

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "submission"
FINAL = SUB / "final"
FIG = FINAL / "figures"
TITLE = "Quantifying error propagation in simulated agentic actuarial workflows: a reproducible operational-risk testbed"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_doc(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 12)]:
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        if name == "Title":
            p_pr = s.element.get_or_add_pPr()
            border = p_pr.find(qn("w:pBdr"))
            if border is not None:
                p_pr.remove(border)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")


def add_field(paragraph, instr: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def add_inline(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def markdown_to_docx(source: str, out: Path, identified: bool):
    doc = Document()
    style_doc(doc)
    lines = source.splitlines()
    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph(title, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if identified:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Adebayo Adetola").bold = True
        p = doc.add_paragraph("Independent Researcher, United Kingdom")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph("Correspondence: adebayoadetola96@yahoo.com")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph("Anonymous manuscript")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in lines[1:]:
        s = line.strip()
        if not s:
            continue
        if not identified:
            s = s.replace("https://github.com/tollyboy88/agentic-actuarial-risk-testbed", "[repository link omitted for review]")
        if s.startswith("## "):
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=2)
        elif s.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), s[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, s)
    doc.add_heading("Figure captions", level=1)
    captions = [
        "Figure 1. Median absolute paired error across workflow stages and control topologies.",
        "Figure 2. Silent-failure rate by workflow topology.",
        "Figure 3. Detection rate by fault class and topology.",
        "Figure 4. Scenario-conditioned annual VaR at 99.5% by topology (logarithmic scale).",
        "Figure 5. Sector VaR at 99.5% under prescribed cross-firm dependence.",
        "Figure 6. Detection sensitivity and 95% world-bootstrap intervals by workflow topology.",
        "Figure 7. Absolute retrospective error by reserving method and CLRD2025 line of business.",
    ]
    for c in captions:
        doc.add_paragraph(c)
    doc.core_properties.title = title
    doc.core_properties.subject = "Agentic actuarial workflow risk simulation"
    doc.core_properties.keywords = "actuarial AI, operational risk, simulation"
    doc.core_properties.author = "Adebayo Adetola" if identified else ""
    doc.core_properties.last_modified_by = "" if not identified else "Adebayo Adetola"
    doc.save(out)


def add_table(doc, title, columns, rows, add_break=True):
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = True
    for i, col in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = col
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if add_break:
        doc.add_page_break()


def make_tables(out: Path):
    capital = pd.read_csv(ROOT / "outputs/scaled/tables/capital.csv")
    summary = pd.read_csv(ROOT / "outputs/scaled/tables/fault_topology_summary.csv")
    systemic = pd.read_csv(ROOT / "outputs/scaled/tables/systemic_stress.csv")
    reviewer = ROOT / "outputs/reviewer_revision/tables"
    detection = pd.read_csv(reviewer / "detection_sensitivity.csv")
    validation = pd.read_csv(reviewer / "clrd_retrospective_validation.csv")
    ablation = pd.read_csv(reviewer / "control_ablation.csv")
    ranking = pd.read_csv(reviewer / "ranking_stability.csv")
    topo = summary.groupby("topology", as_index=False).apply(
        lambda g: pd.Series({"runs": int(g.runs.sum()), "detect": (g.detection_rate*g.runs).sum()/g.runs.sum(), "silent": (g.silent_failure_rate*g.runs).sum()/g.runs.sum()}),
        include_groups=False,
    ).reset_index(drop=True)
    doc = Document()
    style_doc(doc)
    doc.add_paragraph("Tables", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    rows = [[r.topology, int(r.runs), f"{r.detect:.1%}", f"{r.silent:.1%}"] for r in topo.itertuples()]
    add_table(doc, "Table 1. Detection and silent failure by topology", ["Topology", "Fault runs", "Detection", "Silent failure"], rows)
    rows = [[r.topology, f"{r.mean_annual_loss:,.0f}", f"{r.var_995:,.0f}", f"{r.es_995:,.0f}", f"{r.annual_control_cost:,.2f}"] for r in capital.itertuples()]
    add_table(doc, "Table 2. Annual loss and control-cost scenarios (GBP)", ["Topology", "Mean", "VaR 99.5%", "ES 99.5%", "Control cost"], rows)
    med = systemic.groupby("rho").var_995_sector.median().reset_index()
    rows = [[f"{r.rho:.1f}", f"{r.var_995_sector:,.0f}"] for r in med.itertuples()]
    add_table(doc, "Table 3. Median sector VaR 99.5% by prescribed dependence", ["Correlation", "Median sector VaR (GBP)"], rows)
    rows = [
        ["F1", "Data/unit corruption", "Currency or magnitude error; includes ×1,000 stress"],
        ["F2", "Semantic mapping", "Incorrect peril, segment, or field interpretation"],
        ["F3", "Tool arguments", "Correct tool called with an incorrect parameter"],
        ["F4", "Reasoning/judgement", "Unsupported actuarial selection or decision"],
        ["F5", "Hand-off loss", "Required state omitted between agents or stages"],
        ["F6", "Context", "Stale or conflicting context used"],
        ["F7", "Adversarial", "Indirect prompt injection from untrusted content"],
    ]
    add_table(doc, "Table 4. Fault taxonomy", ["Code", "Class", "Operational interpretation"], rows)
    compact = validation.loc[validation.method.isin(["latest_paid", "chain_ladder", "bornhuetter_ferguson"])]
    rows = [[r.lob, r.method.replace("_", " ").title(), f"{r.absolute_percentage_error:.1%}",
             "—" if pd.isna(r.error_ci_low) else f"{r.error_ci_low:.1%} to {r.error_ci_high:.1%}"] for r in compact.itertuples()]
    add_table(doc, "Table 5. CLRD2025 retrospective validation", ["Line", "Method", "Absolute error", "Bootstrap error interval"], rows)
    selected = detection.loc[detection.topology.isin(["T1_LINEAR", "T2_VALIDATOR", "T3_SUPERVISOR", "T5_HUMAN_K2"])]
    rows = [[f"{r.detection_scale:.1f}", r.topology, f"{r.detection_rate:.1%}",
             f"{r.unresolved_failure_rate:.1%}", f"{r.unresolved_ci_low:.1%} to {r.unresolved_ci_high:.1%}"] for r in selected.itertuples()]
    add_table(doc, "Table 6. Detection-probability sensitivity", ["Scale", "Topology", "Detection", "Unresolved", "95% interval"], rows)
    rows = [[r.scenario.replace("_", " "), r.topology, f"{r.detection_rate:.1%}", f"{r.unresolved_failure_rate:.1%}"] for r in ablation.itertuples()]
    add_table(doc, "Table 7. Control ablation", ["Scenario", "Topology", "Detection", "Unresolved failure"], rows)
    rows = [[r.topology, f"{r.share_best:.1%}"] for r in ranking.itertuples()]
    add_table(doc, "Table 8. Lowest-VaR ranking across 135 sensitivity combinations", ["Topology", "Share ranked first"], rows, add_break=False)
    doc.save(out)


def make_title_page(out: Path):
    doc = Document(); style_doc(doc)
    doc.add_paragraph(TITLE, style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in ["Adebayo Adetola", "Independent Researcher, United Kingdom", "Email: adebayoadetola96@yahoo.com", "GitHub: https://github.com/tollyboy88/agentic-actuarial-risk-testbed"]:
        p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Author declarations", level=1)
    doc.add_paragraph("Funding: No external funding was used in the development of this independent research prototype.")
    doc.add_paragraph("Competing interests: The author declares no competing interests.")
    doc.add_paragraph("Author contribution (CRediT): Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Visualization, Writing—original draft, Writing—review and editing: Adebayo Adetola.")
    doc.add_paragraph("ORCID: Not supplied. Add before submission if available.")
    doc.add_paragraph("Telephone: Not supplied; enter in the submission portal if requested.")
    doc.save(out)


def make_cover(out: Path):
    doc = Document(); style_doc(doc)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.add_paragraph("Cover letter", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Editor-in-Chief\nScandinavian Actuarial Journal")
    doc.add_paragraph("Dear Editor,")
    doc.add_paragraph(f"Please consider the manuscript “{TITLE}” as an Original Article.")
    doc.add_paragraph("The manuscript introduces an explicitly simulated, reproducible testbed that traces controlled faults through a six-stage actuarial reserving workflow and translates residual effects into scenario-conditioned operational-risk distributions. The revision adds a 135-combination sensitivity grid, world-bootstrap and Monte Carlo intervals, six-line CLRD2025 retrospective validation, 750 public-data fault injections, and nine control-ablation runs. These additions test robustness while separating methodological demonstration from insurer-specific calibration and from empirical claims about deployed language-model agents.")
    doc.add_paragraph("The work fits the journal’s interest in actuarial methods with practical application. The source code, configurations, public-data provenance, and non-restricted derived results are available at https://github.com/tollyboy88/agentic-actuarial-risk-testbed. No confidential insurer data are used. All references in the manuscript are cited in the text.")
    doc.add_paragraph("This manuscript is original, is not under consideration elsewhere, and has not been published previously. The author declares no competing interests and no external funding. OpenAI Codex assisted with software implementation, literature discovery, language editing, and preparation of submission files; the author retains responsibility for the submitted work.")
    doc.add_paragraph("Thank you for your consideration.\n\nSincerely,\nAdebayo Adetola\nIndependent Researcher, United Kingdom\nadebayoadetola96@yahoo.com")
    doc.save(out)


def make_supplement(out: Path):
    doc = Document(); style_doc(doc)
    doc.add_paragraph("Supplementary material: Agentic Actuarial Testbed", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    sections = {
        "S1. Reproducibility protocol": "Install Python 3.11 or newer, install the package and test dependencies, and run `aat-sim all` with the configuration file `configs/scaled.yaml`. The seed and all resolved parameters are written to the output directory. The full profile is computationally larger and is intended for expanded replication.",
        "S2. Synthetic-world design": "Worlds contain policy exposure, premium, claim ultimate, incremental paid development, report and payment dates, currency, peril, and ambient duplicates. Ground truth is retained before observation defects. No synthetic record represents a real policyholder.",
        "S3. Pairing and replay": "Every injected-fault trajectory is matched to a same-world, same-seed, same-topology no-fault trajectory. A detected fault restores the last clean state and re-executes downstream deterministic stages. Event and snapshot identifiers preserve this relationship.",
        "S4. Capital parameters": "The base screening results use 10,000 annual simulations per topology, configured Poisson event frequencies, a 5% reserve-movement-to-economic-loss conversion, and GBP 5,000 remediation cost. The robustness grid uses 3,000 annual simulations for every combination of detection scale (0.6–1.4), occurrence multiplier (0.5, 1, 2), conversion ratio (1%, 5%, 10%), and remediation cost (GBP 1,000, GBP 5,000, GBP 10,000), with 250 Monte Carlo VaR resamples.",
        "S5. Validation gates": "Completed gates include deterministic world reproduction, six complete finite stage snapshots, visible fault and repair events, Parquet and SQLite persistence, successful aggregation, and 1,024 completed screening runs without hard execution failure.",
        "S6. Public-data validation": "CLRD2025 is evaluated at a calendar-2007 cutoff against paid loss at development lag 10. Latest paid, Chain Ladder, the Mack-identical Chain Ladder point estimate, and Bornhuetter–Ferguson are compared across six lines. One hundred insurer-cluster resamples per line produce retrospective error intervals. A separate 750-case audit injects five deterministic data faults.",
        "S7. Control ablation": "Nine runs remove schema/lineage checks, prompt sanitisation, typed-stage validation, replay/recovery, supervisor review, or human checkpoints. Unresolved paired error is the primary outcome so detection without repair is not counted as risk removal.",
        "S8. Interpretation limits": "The simulator does not contain a live language model and does not estimate production fault frequency, actual human-review effectiveness, firm-specific economic conversion, or cross-firm correlation. CLRD2025 validates reserving and audit mechanics only. The output is suitable for methodological testing and scenario discussion, not a stand-alone regulatory capital estimate.",
        "S9. Recommended extension": "A publication-scale study should increase public-data bootstrap resamples, expand seeded worlds, vary fault magnitudes and tail thresholds, compare Poisson and contagion frequencies, and run preregistered experiments across multiple open and proprietary agent models using frozen prompts and tool traces.",
    }
    for h, text in sections.items():
        doc.add_heading(h, level=1); doc.add_paragraph(text)
    doc.add_heading("S10. File map", level=1)
    for item in ["configs/: experiment assumptions and seeds", "src/aat/: simulation, robustness, and capital engine", "outputs/scaled/: base screening results", "outputs/reviewer_revision/: public-data validation, sensitivity, intervals, and ablations", "dashboard/: interactive result viewer", "tests/: deterministic and end-to-end tests"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(out)


def prepare_figures():
    FIG.mkdir(parents=True, exist_ok=True)
    sources = sorted((ROOT / "outputs/scaled/figures").glob("fig*.png"))
    for i, src in enumerate(sources, 1):
        with Image.open(src) as im:
            im = im.convert("RGB")
            if im.width < 3000:
                scale = 3000 / im.width
                im = im.resize((3000, int(im.height * scale)), Image.Resampling.LANCZOS)
            target = str(FIG / f"Figure_{i}.png")
            if os.name == "nt":
                target = "\\\\?\\" + target
            im.save(target, dpi=(600, 600), optimize=True)

    if plt is None:
        return

    tables = ROOT / "outputs/reviewer_revision/tables"
    detection = pd.read_csv(tables / "detection_sensitivity.csv")
    keep = ["T1_LINEAR", "T2_VALIDATOR", "T3_SUPERVISOR", "T4_CRITIC", "T5_HUMAN_K2"]
    fig, ax = plt.subplots(figsize=(10, 6))
    for topology in keep:
        g = detection.loc[detection.topology == topology].sort_values("detection_scale")
        ax.plot(g.detection_scale, g.detection_rate, marker="o", linewidth=2, label=topology)
        ax.fill_between(g.detection_scale, g.detection_ci_low, g.detection_ci_high, alpha=0.12)
    ax.set(xlabel="Detection-probability scale", ylabel="Detection rate", ylim=(0, 1.05))
    ax.grid(alpha=0.25); ax.legend(ncol=2, frameon=False); fig.tight_layout()
    fig.savefig(FIG / "Figure_6.png", dpi=300); plt.close(fig)

    validation = pd.read_csv(tables / "clrd_retrospective_validation.csv")
    validation = validation.loc[validation.method.isin(["latest_paid", "chain_ladder", "bornhuetter_ferguson"])].copy()
    pivot = validation.pivot(index="lob", columns="method", values="absolute_percentage_error")
    pivot = pivot[["latest_paid", "chain_ladder", "bornhuetter_ferguson"]]
    ax = pivot.rename(columns={"latest_paid": "Latest paid", "chain_ladder": "Chain Ladder", "bornhuetter_ferguson": "Bornhuetter–Ferguson"}).plot.bar(figsize=(10, 6))
    ax.set(xlabel="CLRD2025 line of business", ylabel="Absolute retrospective error")
    ax.yaxis.set_major_formatter(lambda x, pos: f"{x:.0%}")
    ax.grid(axis="y", alpha=0.25); ax.legend(frameon=False); ax.figure.tight_layout()
    ax.figure.savefig(FIG / "Figure_7.png", dpi=300); plt.close(ax.figure)


def main():
    FINAL.mkdir(parents=True, exist_ok=True)
    source = (SUB / "manuscript_source.md").read_text(encoding="utf-8")
    markdown_to_docx(source, FINAL / "01_manuscript_with_author_details.docx", identified=True)
    markdown_to_docx(source, FINAL / "02_anonymous_manuscript.docx", identified=False)
    make_title_page(FINAL / "03_author_details_title_page.docx")
    make_cover(FINAL / "04_cover_letter.docx")
    make_supplement(FINAL / "05_supplementary_material.docx")
    make_tables(FINAL / "06_tables.docx")
    prepare_figures()


if __name__ == "__main__":
    main()
