from __future__ import annotations

import re
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
PAPERS = ROOT / "papers"
DOI = "https://doi.org/10.5281/zenodo.22821051"
REPO = "https://github.com/tollyboy88/agentic-actuarial-risk-testbed"
PAPER1_TITLE = "A Reproducible Fault-Injection Testbed for Error Propagation and Control Evaluation in Agentic Actuarial Reserving Workflows"
PAPER2_TITLE = "Bayesian Multistate Modelling of Error Propagation and Control Placement in Automated Actuarial Workflows"


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = instruction
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); run._r.extend([begin, text, end])


def style_document(doc: Document, *, double: bool = True, uk: bool = False) -> None:
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
        footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_field(footer, "PAGE")
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2 if double else 1.15; normal.paragraph_format.space_after = Pt(0 if double else 6)
    for name, size in [("Title",16),("Heading 1",14),("Heading 2",12),("Heading 3",12)]:
        style = doc.styles[name]; style.font.name = "Times New Roman"; style.font.size = Pt(size); style.font.bold = True


def add_inline(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.*?\*\*)", text):
        run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
        if part.startswith("**") and part.endswith("**"): run.bold = True


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr(); element = OxmlElement("w:tblHeader"); element.set(qn("w:val"), "true"); tr_pr.append(element)


def add_table(doc: Document, title: str, frame: pd.DataFrame, formats: dict[str, str] | None = None) -> None:
    doc.add_paragraph(title).runs[0].bold = True
    table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Table Grid"; set_repeat_table_header(table.rows[0])
    for i, column in enumerate(frame.columns):
        table.rows[0].cells[i].text = str(column); table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for i, column in enumerate(frame.columns):
            value = row[column]
            if formats and column in formats and pd.notna(value): value = formats[column].format(value)
            cells[i].text = "—" if pd.isna(value) else str(value)
    doc.add_paragraph()


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.3) -> None:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def author_block(doc: Document, title: str) -> None:
    p = doc.add_paragraph(title, style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        "Adebayo Aliu Adetola¹ and Olugbenga Akinade²",
        "¹ Independent Researcher; School of Computing, Teesside University, Middlesbrough, United Kingdom",
        "² School of Computing, Teesside University, Middlesbrough, United Kingdom",
        "Corresponding author: Adebayo Aliu Adetola, adebayoadetola96@yahoo.com",
        "ORCID: 0009-0002-1995-6400; 0000-0003-3950-3775",
    ]
    for text in lines:
        p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def anonymise_text(text: str) -> str:
    text = text.replace(REPO, "[repository link withheld during double-anonymized review]")
    text = text.replace(DOI, "[data DOI supplied in the non-anonymous submission metadata]")
    text = re.sub(r"\*\*Author contributions\.\*\*.*?(?=\n\n)", "", text, flags=re.S)
    text = re.sub(r"\*\*Funding\.\*\*.*?(?=\n\n)", "", text, flags=re.S)
    text = re.sub(r"\*\*Competing interests\.\*\*.*?(?=\n\n)", "", text, flags=re.S)
    text = re.sub(r"\*\*AI-use disclosure\.\*\*.*?(?=\n\n)", "", text, flags=re.S)
    text = re.sub(r"\*\*Generative AI statement\.\*\*.*?(?=\n\n)", "", text, flags=re.S)
    text = text.replace("Adetola and Akinade 2026", "Anonymous software release 2026")
    text = text.replace("Adetola, A. A., and O. Akinade. 2026.", "Anonymous. 2026.")
    return text


def markdown_doc(source: str, path: Path, title: str, *, identified: bool, double: bool, journal: str, insert_materials=None) -> None:
    doc = Document(); style_document(doc, double=double, uk=(journal == "AAS"))
    if identified: author_block(doc, title)
    else:
        p=doc.add_paragraph(title, style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p=doc.add_paragraph("Anonymous manuscript"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        source = anonymise_text(source)
    lines = source.splitlines(); first = True
    inserted = False
    for line in lines:
        text = line.strip()
        if first and text.startswith("# "): first=False; continue
        if not text: continue
        if text == "## References" and insert_materials and not inserted:
            insert_materials(doc); inserted=True
        if text.startswith("## "): doc.add_heading(text[3:], level=1)
        elif text.startswith("### "): doc.add_heading(text[4:], level=2)
        elif text.startswith("#### "): doc.add_heading(text[5:], level=3)
        elif text.startswith("- "): add_inline(doc.add_paragraph(style="List Bullet"), text[2:])
        else: add_inline(doc.add_paragraph(), text)
    doc.core_properties.title = title; doc.core_properties.author = "Adebayo Aliu Adetola; Olugbenga Akinade" if identified else ""
    doc.core_properties.keywords = "actuarial workflow; error propagation; Bayesian multistate model; control evaluation"
    doc.core_properties.last_modified_by = "" if not identified else "Adebayo Aliu Adetola"
    path.parent.mkdir(parents=True, exist_ok=True); doc.save(path)


def paper1_source(aas: bool = False) -> str:
    source = (ROOT / "submission/manuscript_source.md").read_text(encoding="utf-8")
    source = re.sub(r"^# .*", f"# {PAPER1_TITLE}", source, count=1)
    abstract = (
        "Automated actuarial workflows can propagate a local data, tool, judgement, hand-off, context or adversarial fault into a material reserve misstatement. This article introduces a reproducible fault-injection testbed that executes six auditable reserving stages, pairs each faulty trajectory with a same-world clean counterfactual, and compares automated and human control architectures with restore-and-replay recovery. The fixed screening design contains 1,024 runs across eight synthetic insurer worlds. Per-stage validation reduced silent failure from 81.7% in the unreviewed chain to 2.5%, but a 135-scenario sensitivity grid showed that the lowest-risk control was not invariant: the validator ranked first in 40.0% of scenarios, a two-checkpoint human design in 31.1%, and the supervisor in 28.9%. Removing replay increased unresolved validator failures from 0.8% to 40.0%. A retrospective CLRD2025 experiment across six lines of business produced mean absolute errors of 6.2% for Chain Ladder, 8.9% for Bornhuetter–Ferguson and 26.3% for latest paid. Adding a local scale-consistency control increased detection of injected ×1,000 unit faults from 22.7% to 74.0%. The contribution is an applied experimental and audit framework for control design, not an estimate of production language-model failure frequency or insurer regulatory capital. Code, configurations, tests and unrestricted reproduction materials are publicly archived."
    )
    if aas:
        abstract = ("We introduce an open-source Python testbed for controlled fault injection and control evaluation in multi-stage actuarial reserving workflows. The software generates seeded known-truth portfolios or accepts public loss-development data, executes six auditable stages, injects seven fault classes, and compares automated and human controls using paired clean counterfactuals. Detected faults can be restored and replayed while event ledgers preserve lineage. Reproducible examples include 1,024 synthetic runs, retrospective validation on six CLRD2025 lines, public-data fault audits, detection sensitivity and component ablation. Per-stage validation reduced silent failure from 81.7% to 2.5%; removing replay raised unresolved validator failures from 0.8% to 40.0%. A new local scale check increased detection of ×1,000 unit faults from 22.7% to 74.0%. The package reports propagation, unresolved material error, control cost and scenario-conditioned tail-risk measures. It is a reusable experimental and audit toolbox, not an estimate of production agent-failure frequencies or regulatory capital.")
    source = re.sub(r"(?s)(## Abstract\n\n).*?(\n\n\*\*Keywords:)", rf"\1{abstract}\2", source, count=1)
    source = re.sub(r"\*\*Keywords:\*\*.*", "**Keywords:** actuarial artificial intelligence; fault injection; claims reserving; model risk; workflow controls", source, count=1)
    source = source.replace("**Practitioner summary:**", "**Practical applications:**")
    source = source.replace(
        "The deterministic audit detected 84.5% of 750 injected public-data faults. It detected all duplicate-key, negative-paid, development-mismatch, and missing-cell faults, but only 22.7% of multiplication-by-1,000 faults. A large value is not necessarily anomalous in a heterogeneous insurer portfolio, so a generic extreme-increment rule cannot replace unit metadata, reconciliation totals, and insurer-specific thresholds.",
        "The legacy deterministic audit detected all duplicate-key, negative-paid, development-mismatch and missing-cell faults, but only 22.7% of multiplication-by-1,000 faults. A new local scale-consistency check compares the target cumulative payment with the median positive peer development value in the same accident year. With a deliberately conservative 100-fold threshold, unit-fault detection increased to 74.0%. The residual cases lacked a stable positive peer or did not cross the threshold, so unit metadata and reconciliation totals remain necessary."
    )
    source = source.replace("**Funding.** No external funding was used in the development of this independent research prototype.", "**Funding.** This work received no specific grant from any funding agency, commercial or not-for-profit sectors.")
    source = source.replace("**Competing interests.** The author declares no competing interests.", "**Competing interests.** The authors declare no competing interests.\n\n**Author contributions.** Adebayo Aliu Adetola: conceptualisation, methodology, software, formal analysis, investigation, data curation, visualisation, writing—original draft, and writing—review and editing. Olugbenga Akinade: supervision, methodology, validation, and writing—review and editing.")
    source = source.replace("Configurations, source provenance, derived results, and reproduction instructions are available at https://github.com/tollyboy88/agentic-actuarial-risk-testbed.", f"The dataset archive is available at {DOI}; configurations, source provenance, derived results and reproduction instructions are available at {REPO}.")
    source = source.replace("**Code availability.** Source code is available under the MIT License at https://github.com/tollyboy88/agentic-actuarial-risk-testbed.", f"**Code availability.** Source code version 1.0.0 is available at {REPO} under the MIT License.")
    source = source.replace("**Generative AI statement.**", "**AI-use disclosure.**")
    source = source.replace("## 1. Introduction", "**JEL classification:** C15; C63; G22\n\n## 1. Introduction")
    source = source.replace("## References", "This manuscript cites Figure 1 for the complete architecture; Figures 2–7 for propagation, detection, public-data validation, sensitivity, ablation and ranking stability; and Tables 1–7 for the corresponding definitions and numerical results.\n\n## References")
    if aas:
        source = source.replace("## 4. Results", "### 3.9 Implementation and user interface\n\nThe Python package exposes a command-line interface, YAML configuration, Parquet/CSV tables, SQLite event trajectories and a local dashboard. `scripts/reproduce_naaj.py` regenerates the fixed applied study; `scripts/reproduce_saj.py` builds the probabilistic extension. A run manifest stores SHA-256 hashes, and `REPRODUCIBILITY.md` gives installation and resource instructions.\n\n### 3.10 Unit Testing and Validation\n\nThe test suite checks deterministic seeds, unique run identifiers, six stage records per run, transition probabilities, analytic state recursion, known-parameter recovery and required publication outputs. Nine tests pass in the frozen environment. The smoke profile supports rapid installation checks; the scaled profile recreates the reported 1,024-run study.\n\n## 4. Results")
        source = source.replace("American", "American")
    return source


def make_architecture_figures() -> tuple[Path, Path]:
    asset = PAPERS / "assets"; asset.mkdir(parents=True, exist_ok=True)
    p1 = asset / "paper1_architecture.png"; fig, ax = plt.subplots(figsize=(12,5)); ax.axis("off")
    boxes = [(0.03,"Synthetic / CLRD2025\ndata"),(0.21,"Six-stage reserving\nworkflow"),(0.40,"Seven controlled\nfault classes"),(0.59,"Alternative controls\nand replay"),(0.78,"Paired error, audit\nand scenario risk")]
    for x,label in boxes:
        ax.add_patch(plt.Rectangle((x,.35),.16,.3,facecolor="#e8f1fa",edgecolor="#1f4e79",lw=1.5)); ax.text(x+.08,.5,label,ha="center",va="center",fontsize=10)
    for x,_ in boxes[:-1]: ax.annotate("",xy=(x+.20,.5),xytext=(x+.16,.5),arrowprops={"arrowstyle":"->","lw":1.5})
    ax.text(.5,.14,"Every treatment trajectory is paired with a same-world clean counterfactual; detected faults restore the last clean checkpoint and replay downstream stages.",ha="center",fontsize=9)
    fig.tight_layout(); fig.savefig(p1,dpi=300,bbox_inches="tight"); plt.close(fig)
    p2 = asset / "paper2_state_process.png"; fig,ax=plt.subplots(figsize=(11,5)); ax.axis("off")
    nodes={"C":(.10,.52,"Clean"),"E":(.38,.52,"Active error"),"R":(.70,.72,"Recovered"),"H":(.70,.32,"Hard failure")}
    for key,(x,y,label) in nodes.items(): ax.add_patch(plt.Rectangle((x-.08,y-.07),.16,.14,facecolor="#f6f8fb",edgecolor="#244a73",lw=2)); ax.text(x,y,f"{key}: {label}",ha="center",va="center")
    ax.annotate("",xy=(.30,.52),xytext=(.18,.52),arrowprops={"arrowstyle":"->","lw":1.5}); ax.text(.24,.56,"fault",ha="center")
    ax.annotate("",xy=(.62,.70),xytext=(.46,.56),arrowprops={"arrowstyle":"->","lw":1.5}); ax.text(.54,.68,"detect + replay",ha="center")
    ax.annotate("",xy=(.62,.34),xytext=(.46,.48),arrowprops={"arrowstyle":"->","lw":1.5}); ax.text(.54,.35,"execution failure",ha="center")
    ax.annotate("persists",xy=(.36,.59),xytext=(.28,.78),arrowprops={"arrowstyle":"->","connectionstyle":"arc3,rad=.55"},ha="center")
    ax.text(.5,.10,"Hierarchical transition posterior → Markov-additive error magnitude → posterior loss → constrained control optimiser",ha="center",fontsize=10,weight="bold")
    fig.tight_layout(); fig.savefig(p2,dpi=300,bbox_inches="tight"); plt.close(fig); return p1,p2


def paper1_materials(doc: Document) -> None:
    assets = PAPERS / "assets"; out = ROOT / "outputs"
    doc.add_heading("Integrated tables and figures", level=1)
    add_figure(doc, assets/"paper1_architecture.png", "Figure 1. Study architecture from data and controlled fault injection to paired outcomes and scenario-conditioned risk.")
    figures = [
        (out/"scaled/figures/fig1_error_propagation.png","Figure 2. Stage-wise error propagation by topology; amplification is capped at 100 for display."),
        (out/"scaled/figures/fig3_detection_heatmap.png","Figure 3. Experimental detection rate by fault class and injection stage."),
        (PAPERS/"assets/clrd_validation.png","Figure 4. CLRD2025 retrospective absolute error by line and reserving method."),
        (PAPERS/"assets/detection_sensitivity.png","Figure 5. Detection sensitivity with world-bootstrap intervals."),
        (PAPERS/"assets/ablation.png","Figure 6. Unresolved failure in component-ablation experiments."),
        (PAPERS/"assets/ranking.png","Figure 7. Share of sensitivity scenarios in which each topology has the lowest VaR."),
    ]
    for path,caption in figures: add_figure(doc,path,caption)
    doc.add_page_break()
    summary=pd.read_csv(out/"scaled/tables/fault_topology_summary.csv"); runs=pd.read_parquet(out/"scaled/runs.parquet")
    topo=runs[runs.fault_type!="CONTROL"].groupby("topology",as_index=False).agg(Fault_runs=("run_id","size"),Detection=("detected","mean"),Silent_failure=("silent_failure","mean"))
    faults=pd.DataFrame([("F1","Data/unit corruption"),("F2","Semantic mapping"),("F3","Tool arguments"),("F4","Reasoning/judgement"),("F5","Hand-off loss"),("F6","Context"),("F7","Adversarial")],columns=["Code","Class"])
    add_table(doc,"Table 1. Fault taxonomy",faults)
    topologies=pd.DataFrame([("T1","Linear chain","No dedicated review"),("T2","Validator","Typed validation after each stage"),("T3","Supervisor","Cross-stage review and replay"),("T4","Critic","Model/selection-focused critique"),("T5","Human checkpoints","Zero to three checkpoints")],columns=["Family","Name","Control placement"])
    add_table(doc,"Table 2. Control topology definitions",topologies)
    add_table(doc,"Table 3. Primary topology performance",topo,{"Detection":"{:.1%}","Silent_failure":"{:.1%}"})
    clrd=pd.read_csv(out/"reviewer_revision/tables/clrd_retrospective_validation.csv"); clrd=clrd[clrd.method.isin(["latest_paid","chain_ladder","bornhuetter_ferguson"])][["lob","method","absolute_percentage_error"]]
    add_table(doc,"Table 4. CLRD2025 retrospective validation",clrd,{"absolute_percentage_error":"{:.1%}"})
    audit=pd.read_csv(out/"reviewer_revision/tables/clrd_fault_audit.csv").groupby("fault",as_index=False).agg(Legacy_detection=("detected_legacy","mean"),Enhanced_detection=("detected","mean"))
    add_table(doc,"Table 5. Public-data fault audit",audit,{"Legacy_detection":"{:.1%}","Enhanced_detection":"{:.1%}"})
    ranking=pd.read_csv(out/"reviewer_revision/tables/ranking_stability.csv"); add_table(doc,"Table 6. Ranking stability over 135 scenarios",ranking,{"share_best":"{:.1%}"})
    ablation=pd.read_csv(out/"reviewer_revision/tables/control_ablation.csv")[["scenario","topology","detection_rate","unresolved_failure_rate"]]
    ablation["scenario"]=ablation["scenario"].str.replace("_"," ")
    ablation.columns=["Scenario","Topology","Detection","Unresolved"]
    add_table(doc,"Table 7. Control ablation",ablation,{"Detection":"{:.1%}","Unresolved":"{:.1%}"})


def make_secondary_figures() -> None:
    assets=PAPERS/"assets"; assets.mkdir(parents=True,exist_ok=True); rev=ROOT/"outputs/reviewer_revision/tables"
    clrd=pd.read_csv(rev/"clrd_retrospective_validation.csv"); clrd=clrd[clrd.method.isin(["latest_paid","chain_ladder","bornhuetter_ferguson"])]
    pivot=clrd.pivot(index="lob",columns="method",values="absolute_percentage_error"); ax=pivot.plot.bar(figsize=(9,5)); ax.set_ylabel("Absolute retrospective error"); ax.figure.tight_layout(); ax.figure.savefig(assets/"clrd_validation.png",dpi=300); plt.close(ax.figure)
    det=pd.read_csv(rev/"detection_sensitivity.csv"); fig,ax=plt.subplots(figsize=(9,5))
    for topology in ["T1_LINEAR","T2_VALIDATOR","T3_SUPERVISOR","T5_HUMAN_K2"]:
        g=det[det.topology==topology].sort_values("detection_scale"); ax.plot(g.detection_scale,g.detection_rate,marker="o",label=topology); ax.fill_between(g.detection_scale,g.detection_ci_low,g.detection_ci_high,alpha=.12)
    ax.set(xlabel="Detection scale",ylabel="Detection rate",ylim=(0,1.05)); ax.legend(); fig.tight_layout(); fig.savefig(assets/"detection_sensitivity.png",dpi=300); plt.close(fig)
    abl=pd.read_csv(rev/"control_ablation.csv"); fig,ax=plt.subplots(figsize=(10,5)); ax.barh(abl.scenario,abl.unresolved_failure_rate); ax.set_xlabel("Unresolved failure rate"); fig.tight_layout(); fig.savefig(assets/"ablation.png",dpi=300); plt.close(fig)
    rank=pd.read_csv(rev/"ranking_stability.csv"); fig,ax=plt.subplots(figsize=(7,4)); ax.bar(rank.topology,rank.share_best); ax.set_ylabel("Share ranked first"); ax.tick_params(axis="x",rotation=25); fig.tight_layout(); fig.savefig(assets/"ranking.png",dpi=300); plt.close(fig)


def paper2_materials(doc: Document) -> None:
    assets=PAPERS/"assets"; out=ROOT/"outputs/saj_methodology"; doc.add_heading("Integrated tables and figures",level=1)
    add_figure(doc,assets/"paper2_state_process.png","Figure 1. Four-state workflow process and its connection to error magnitude, posterior risk and control choice.")
    add_figure(doc,out/"figures/figure2_transition_intervals.png","Figure 2. Posterior recovery probabilities by control topology, averaged over observed cells.")
    add_figure(doc,out/"figures/figure3_parameter_recovery.png","Figure 3. Empirical coverage of nominal 90% intervals in the known-parameter study.")
    scores=pd.read_csv(out/"tables/held_out_model_comparison.csv"); fig,ax=plt.subplots(figsize=(6,4)); x=np.arange(len(scores)); ax.bar(x-.18,scores.log_score,.36,label="Log score"); ax.bar(x+.18,scores.brier_score,.36,label="Brier score"); ax.set_xticks(x,scores.model); ax.legend(); fig.tight_layout(); score_path=assets/"model_scores.png"; fig.savefig(score_path,dpi=300); plt.close(fig)
    add_figure(doc,score_path,"Figure 4. Held-out-world probabilistic scores; lower values are better.")
    add_figure(doc,out/"figures/figure5_pareto_frontier.png","Figure 5. Cost-risk Pareto frontier for 64 stage-control placements.")
    states=pd.DataFrame([("C","Clean execution"),("E","Active material error"),("R","Detected and recovered"),("H","Hard execution failure")],columns=["State","Definition"]); add_table(doc,"Table 1. State definitions",states)
    priors=pd.DataFrame([("Weak",.5,"Minimal pooling in recovery study"),("Main",8,"Application baseline"),("Sceptical",12,"Stronger shrinkage sensitivity")],columns=["Prior","Transition strength κ","Use"]); add_table(doc,"Table 2. Prior specifications",priors)
    score_table=scores[["model","n","log_score","brier_score"]].copy(); score_table.columns=["Model","N","Log score","Brier score"]
    add_table(doc,"Table 3. Held-out model comparison",score_table,{"Log score":"{:.3f}","Brier score":"{:.3f}"})
    recovery=pd.read_csv(out/"tables/parameter_recovery_summary.csv"); recovery["topology"]=recovery["topology"].replace({"T1_LINEAR":"Linear","T2_VALIDATOR":"Validator"}); recovery.columns=["Topology","Bias","RMSE","Coverage 90%","Coverage 95%"]
    add_table(doc,"Table 4. Simulation recovery",recovery,{"Bias":"{:.3f}","RMSE":"{:.3f}","Coverage 90%":"{:.1%}","Coverage 95%":"{:.1%}"})
    opt=pd.read_csv(out/"tables/control_optimisation.csv",dtype={"placement":str}); opt=opt[opt.feasible].sort_values("objective").head(8)[["placement","controls","annual_cost","posterior_mean_silent_probability","selected"]]; opt["placement"]=opt["placement"].str.zfill(6); opt.columns=["Placement","Controls","Cost","P(silent)","Selected"]
    add_table(doc,"Table 5. Leading budget-feasible control placements",opt,{"P(silent)":"{:.4f}"})


def cover_letter(path: Path, journal: str, title: str, body: list[str]) -> None:
    doc=Document(); style_document(doc,double=False); doc.add_paragraph("Cover letter",style="Title").alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Editor-in-Chief\n{journal}\n\nDear Editor,")
    doc.add_paragraph(f"Please consider our manuscript “{title}”.")
    for paragraph in body: doc.add_paragraph(paragraph)
    doc.add_paragraph("The work is original, is not under consideration elsewhere, and has not been published. The authors declare no competing interests and no specific external funding.")
    doc.add_paragraph("Sincerely,\nAdebayo Aliu Adetola, corresponding author\nadebayoadetola96@yahoo.com\nOn behalf of both authors")
    path.parent.mkdir(parents=True,exist_ok=True); doc.save(path)


def title_page(path: Path, title: str, journal: str) -> None:
    doc=Document(); style_document(doc,double=False); author_block(doc,title); doc.add_heading("Submission details",level=1); doc.add_paragraph(f"Target journal: {journal}"); doc.add_paragraph("Article type: Original research article")
    doc.add_heading("Declarations",level=1); doc.add_paragraph("Funding: This work received no specific grant from any funding agency, commercial or not-for-profit sectors."); doc.add_paragraph("Competing interests: The authors declare no competing interests."); doc.add_paragraph(f"Data DOI: {DOI}"); doc.add_paragraph(f"Code repository: {REPO}"); doc.save(path)


def supplement(path: Path, title: str, sections: list[tuple[str,str]]) -> None:
    doc=Document(); style_document(doc,double=False); doc.add_paragraph(title,style="Title").alignment=WD_ALIGN_PARAGRAPH.CENTER
    for heading,text in sections: doc.add_heading(heading,level=1); doc.add_paragraph(text)
    doc.save(path)


def main() -> None:
    make_architecture_figures(); make_secondary_figures()
    naaj=PAPERS/"paper1_naaj/submission"; aas=PAPERS/"paper1_aas/submission"; saj=PAPERS/"paper2_saj/submission"
    for package in [naaj, aas, saj]: package.mkdir(parents=True, exist_ok=True)
    source=paper1_source(False); (naaj.parent/"manuscript_source.md").write_text(source,encoding="utf-8")
    markdown_doc(source,naaj/"01_manuscript_with_author_details.docx",PAPER1_TITLE,identified=True,double=True,journal="NAAJ",insert_materials=paper1_materials)
    markdown_doc(source,naaj/"02_anonymous_manuscript.docx",PAPER1_TITLE,identified=False,double=True,journal="NAAJ",insert_materials=paper1_materials)
    cover_letter(naaj/"03_cover_letter.docx","North American Actuarial Journal",PAPER1_TITLE,["The article presents an applied, reproducible actuarial fault-injection framework rather than a claim of new probability theory. It pairs controlled faults with clean counterfactuals, compares control and replay designs, validates reserving mechanics on public CLRD2025 triangles, and reports sensitivity, uncertainty and ablation evidence.",f"The data archive is {DOI}; code version 1.0.0 and the one-command reproduction scripts are at {REPO}. We request the standard non-open-access publication route, for which the journal states there is no APC."])
    supplement(naaj/"04_supplementary_material.docx","Supplementary material for Paper 1",[("S1 Reproduction","Use Python 3.11 or newer, install `.[test]`, run `python scripts/reproduce_naaj.py`, and verify the SHA-256 manifest."),("S2 Parameter boundary","Detection, occurrence, economic conversion, remediation cost and dependence are explicit scenario inputs; none is presented as an industry estimate."),("S3 Public-data audit","The CLRD audit contains 750 deterministic injections. The enhanced local scale check raises ×1,000 detection from 22.7% to 74.0%; unit metadata and reconciliation remain recommended."),("S4 Full outputs",f"Machine-readable tables, configurations and provenance are at {REPO}; the open dataset archive is {DOI}.")])
    (naaj/"05_reproducibility_link.txt").write_text(f"Dataset: {DOI}\nCode: {REPO}\nRelease: v1.0.0\nCommands: python scripts/reproduce_naaj.py; python -m pytest\n",encoding="utf-8")

    aas_source=paper1_source(True); (aas.parent/"manuscript_source.md").write_text(aas_source,encoding="utf-8")
    markdown_doc(aas_source,aas/"01_manuscript_AAS.docx",PAPER1_TITLE,identified=True,double=False,journal="AAS",insert_materials=paper1_materials)
    cover_letter(aas/"02_cover_letter_AAS.docx","Annals of Actuarial Science",PAPER1_TITLE,["This fallback submission is positioned for the Actuarial Software stream: it documents a reusable Python toolbox, reproducible synthetic and public-data worked examples, unit testing, event-ledger outputs, a dashboard and comparisons with standard reserving tools.","Before this fallback is submitted, the software release must complete a source-ownership review and be relicensed GPL-3.0-only, as required by the Actuarial Software guidance. The current NAAJ-first release remains MIT; the manuscript should therefore be held until that release step is completed."])
    supplement(aas/"03_replication_readme.docx","Replication readme",[("Environment","Python 3.11–3.13; install with `python -m pip install -e .[test]`."),("Commands","Run `python scripts/reproduce_naaj.py` and `python -m pytest`. The dashboard instructions are in `dashboard/README.md`."),("Stable resources",f"Dataset DOI: {DOI}. Source: {REPO}."),("AAS release gate","Before submission, confirm code ownership, change the repository licence to GPL-3.0-only, tag the release, update CITATION.cff and archive the GPL release with Zenodo.")])
    supplement(aas/"04_supplementary_methods.docx","Supplementary methods",[("Experimental design","The applied study uses 1,024 paired runs over eight known-truth worlds, seven fault classes, six stages and eight topology variants."),("Public validation","CLRD2025 retrospective validation covers six lines and 772 insurer-line records. The data-quality audit applies 750 deterministic injections."),("Sensitivity and ablation","Detection is rescaled from 60% to 140%; economic assumptions form 135 combinations; nine ablations remove individual controls or replay."),("Evidence boundary","The software does not estimate deployed-agent frequency, insurer-specific conversion or regulatory capital.")])

    p2=(PAPERS/"paper2_saj/manuscript_source.md").read_text(encoding="utf-8")
    markdown_doc(p2,saj/"01_main_manuscript_anonymous.docx",PAPER2_TITLE,identified=False,double=True,journal="SAJ",insert_materials=paper2_materials)
    markdown_doc(p2,saj/"02_manuscript_with_author_details.docx",PAPER2_TITLE,identified=True,double=True,journal="SAJ",insert_materials=paper2_materials)
    title_page(saj/"03_title_page.docx",PAPER2_TITLE,"Scandinavian Actuarial Journal")
    cover_letter(saj/"04_cover_letter.docx","Scandinavian Actuarial Journal",PAPER2_TITLE,["This is a new methodological manuscript. An earlier testbed manuscript received a desk rejection that identified insufficient probabilistic/statistical innovation. The present paper has a different research question, new code, new analyses and new results: an estimable hierarchical multistate model, Markov-additive error propagation, formal recursion and monotonicity statements, known-parameter recovery, held-out prediction, and optimal control placement.","The applied software paper is distinct experimental infrastructure. It is disclosed and cited where appropriate. This submission does not duplicate its long methods, tables, figures or prose.",f"Data are archived at {DOI}; reproducible code and fixed outputs are at {REPO}. We request the standard non-open-access route, for which the journal states there is no APC."])
    supplement(saj/"05_supplement_proofs_diagnostics.docx","Supplementary proofs and diagnostics",[("Proof of Proposition 1","For any destination state j, Pr(Z_{s+1}=j)=Σ_i Pr(Z_s=i)P_s(i,j), hence π_{s+1}=π_sP_s. Induction gives π_S=π_0Π_sP_s. Multiplying the E component by the conditional materiality probability gives terminal silent failure."),("Proof of Proposition 2","Couple base and improved paths with the same uniforms. Moving active-error probability mass only from E to absorbing recovery R makes every base recovery also an improved recovery and may recover additional paths. Therefore the improved terminal indicator 1{Z_S=E} is no greater almost surely."),("Finite optimum","The admissible catalogue has 64 elements. A finite real-valued objective attains its minimum; exhaustive enumeration evaluates every element and therefore returns a global minimizer."),("Diagnostics","Known-parameter 90%/95% coverage is 92.5%/96.9%. Held-out hierarchical log/Brier scores are 0.360/0.195 versus 1.076/0.652 globally. Analytic and Monte Carlo state probabilities differ by at most 0.000635."),("Prior sensitivity","Transition concentration values 0.5, 8 and 12 define weak, main and sceptical pooling. Production use should add hyperpriors and full posterior predictive checks when more worlds are observed.")])
    supplement(saj/"06_replication_readme.docx","Paper 2 replication readme",[("Environment","Python 3.11 or newer. Install the repository and test dependencies."),("Command","Run `python scripts/reproduce_saj.py`. Outputs are written to `outputs/saj_methodology/`; then run `python -m pytest`."),("Inputs","Paper 2 consumes the frozen `outputs/scaled/runs.parquet` and `stage_metrics.parquet` produced by Paper 1."),("Outputs","Transition observations, posterior tables, recovery study, held-out scores, recursion validation, control enumeration, figures and SHA-256 run manifest."),("Availability",f"Dataset: {DOI}. Code: {REPO}.")])

    for package in [naaj,aas,saj]:
        figures=package/"figures"; figures.mkdir(exist_ok=True)
        for image in sorted((PAPERS/"assets").glob("*.png")): shutil.copy2(image,figures/image.name)

    paper1_figure_sources = [
        PAPERS/"assets/paper1_architecture.png",
        ROOT/"outputs/scaled/figures/fig1_error_propagation.png",
        ROOT/"outputs/scaled/figures/fig3_detection_heatmap.png",
        PAPERS/"assets/clrd_validation.png", PAPERS/"assets/detection_sensitivity.png",
        PAPERS/"assets/ablation.png", PAPERS/"assets/ranking.png",
    ]
    paper2_figure_sources = [
        PAPERS/"assets/paper2_state_process.png",
        ROOT/"outputs/saj_methodology/figures/figure2_transition_intervals.png",
        ROOT/"outputs/saj_methodology/figures/figure3_parameter_recovery.png",
        PAPERS/"assets/model_scores.png",
        ROOT/"outputs/saj_methodology/figures/figure5_pareto_frontier.png",
    ]
    for package in [naaj, aas]:
        for i, source_path in enumerate(paper1_figure_sources, 1):
            with Image.open(source_path) as image:
                image.convert("RGB").save(package/"figures"/f"Figure_{i}.tiff", dpi=(300,300), compression="tiff_lzw")
    for i, source_path in enumerate(paper2_figure_sources, 1):
        with Image.open(source_path) as image:
            image.convert("RGB").save(saj/"figures"/f"Figure_{i}.tiff", dpi=(300,300), compression="tiff_lzw")


if __name__ == "__main__":
    main()
